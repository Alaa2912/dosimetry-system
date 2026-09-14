import os
from docx import Document
from docx.shared import Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
import openpyxl

BASE_DIR = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
TEMPLATES_DIR = os.path.join(BASE_DIR, "templates")
os.makedirs(TEMPLATES_DIR, exist_ok=True)

def build_excel_templates():
    wb_m = openpyxl.Workbook()
    ws_m = wb_m.active
    ws_m.title = "Monthly Report"
    headers_m = [
        "No", "Facility Name", "Serial Number", "Name", "Started Date",
        "Reception Date (related to the currentMonth)", "Reading Date (related to the current Month)",
        "Report Date", "Month", "Hp(10)in mSv", "Hp(0.07)in mSv", "Accumulated Dose (mSv)", "Remark"
    ]
    ws_m.append(headers_m)
    ws_m.append(["", "", "", "", "", "تاريخ استلام العينة", "تاريخ قراءة العينة (المقاييس)", "", "", "", "", "", ""])
    wb_m.save(os.path.join(TEMPLATES_DIR, "FinalDR_monthly.xlsx"))

    wb_q = openpyxl.Workbook()
    ws_q = wb_q.active
    ws_q.title = "Quarterly Report"
    headers_q = [
        "No", "Facility Name", "Serial Number", "Name", "Started Date",
        "Reception Date (related to the current quarter)", "Reading Date (related to the current quarter)",
        "Report Date", "Quarter", "Hp(10)in mSv", "Hp(0.07)in mSv", "Accumulated Dose (mSv)", "Remark"
    ]
    ws_q.append(headers_q)
    wb_q.save(os.path.join(TEMPLATES_DIR, "FinalDR_Quartar.xlsx"))
    print("Excel templates generated.")

def create_base_docx(period_type: str, sub_title: str, month_names: list, filename: str, is_quarter: bool = False):
    doc = Document()
    
    p = doc.add_paragraph()
    run = p.add_run("PERSONAL MONITORING SERVICE\n")
    run.bold = True
    run.font.size = Pt(14)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    
    hdr_p = doc.add_paragraph()
    hdr_p.add_run("Customer Name: {{ customer_name }}\n")
    hdr_p.add_run("Customer Code: {{ customer_code }}\n")
    hdr_p.add_run("Customer Address: {{ customer_address }}\n")
    hdr_p.add_run("Contact: {{ contact }}\n")
    hdr_p.add_run("Monitoring Period (MP): {{ monitoring_period }}\n")
    hdr_p.add_run("Classification: Medical Field\n")
    hdr_p.add_run("Number of Dosimeters: {{ dosimeter_count }}\n")
    hdr_p.add_run("Agreement Date: {{ agreement_date }}\n")
    hdr_p.add_run("Department: {{ department_name }}\n")
    
    p2 = doc.add_paragraph()
    p2.add_run("METHOD USED:\n").bold = True
    p2.add_run("Purpose: Assessment of occupational exposure due to external sources of radiation. Dosimeter System: OSL\n")
    p2.add_run("Quantity and Unit: Doses are reported in terms of Personal Dose Equivalent Hp(10) and Hp(0.07) in units of milli-Sievert (mSv).\n")
    
    p_reg = doc.add_paragraph()
    p_reg.add_run("(علماً بأن جميع القراءات تقع ضمن الحدود المسموح بها حسب تعليمات حدود الجرعات الإشعاعية الصادرة عن مجلس مفوضي هيئة تنظيم قطاع الطاقة والمعادن لسنة 2015)")
    p_reg.runs[0].font.size = Pt(9)
    p_reg.alignment = WD_ALIGN_PARAGRAPH.CENTER
    
    p_rep = doc.add_paragraph()
    p_rep.add_run("PERSONAL DOSE MONITORING REPORT").bold = True
    p_rep.alignment = WD_ALIGN_PARAGRAPH.CENTER
    
    col_count = 17 if is_quarter else 16
    table = doc.add_table(rows=2, cols=col_count)
    table.style = 'Table Grid'
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    
    hdr_cells = table.rows[0].cells
    idx = 0
    hdr_cells[idx].text = "No."
    idx += 1
    if is_quarter:
        hdr_cells[idx].text = "SN"
        idx += 1
    hdr_cells[idx].text = "Name"
    idx += 1
    hdr_cells[idx].text = "Started Date"
    idx += 1
    hdr_cells[idx].text = "Reception Date"
    idx += 1
    hdr_cells[idx].text = "Reading Date"
    idx += 1
    
    for m in month_names:
        hdr_cells[idx].text = f"({m}) in mSv"
        idx += 2
        
    hdr_cells[idx].text = "Accumulated Dose (mSv)"
    idx += 2
    hdr_cells[idx].text = "Remark"
    
    sub_cells = table.rows[1].cells
    sub_start = 6 if is_quarter else 5
    for i in range(5):
        sub_cells[sub_start + (i*2)].text = "Hp(10)"
        sub_cells[sub_start + (i*2) + 1].text = "Hp(0.07)"
        
    doc.add_paragraph()
    p_ft = doc.add_paragraph()
    p_ft.add_run("Remarks and Abbreviations:\n").bold = True
    p_ft.add_run("NI: Not Issued | ES: Estimated dose | S: Stopped | NR: Not Returned\n")
    p_ft.add_run("Hp(10): Whole body dose | HP(0.07): Skin dose | UA: Under Apron | OA: Over Apron | LLD: Lower Limits of Detector (0.01 mSv)\n")
    p_ft.add_run("Type of Detector: OSL (Optical Stimulated Luminescence)\n")
    p_ft.runs[0].font.size = Pt(8.5)
    
    p_sig = doc.add_paragraph()
    p_sig.add_run("\nMeasured by: ____________________       Approved by: ____________________\nSignature:   ____________________       Signature:   ____________________")
    
    doc.save(os.path.join(TEMPLATES_DIR, filename))
    print(f"Saved {filename}")

def build_all_templates():
    build_excel_templates()
    create_base_docx("Monthly", "Months 1-4", ["January", "February", "March", "April"], "Monthly 1-4.docx")
    create_base_docx("Monthly", "Months 5-8", ["May", "June", "July", "August"], "Monthly 5-8.docx")
    create_base_docx("Monthly", "Months 9-12", ["September", "October", "November", "December"], "Monthly 9-12.docx")
    create_base_docx("Quarterly", "Quarterly", ["1st Quarter", "2nd Quarter", "3rd Quarter", "4th Quarter"], "PERSONAL MONITORING SERVICE.docx", is_quarter=True)

if __name__ == "__main__":
    build_all_templates()
